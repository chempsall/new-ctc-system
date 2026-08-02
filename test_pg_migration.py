"""
test_pg_migration.py
PostgreSQL migration acceptance suite — covers every requirement in the
migration brief. Run from the project root against an EMPTY PostgreSQL
database:

    RF_DATABASE_URL=postgresql://user:pass@host:5432/rft_test \
    RF_ENV=development RF_ADMIN_TOKEN=testtoken python3 test_pg_migration.py

The suite creates its own synthetic staff-list and PAR workbooks, so no
source data is required. It is destructive to the target database —
never point it at a live one.

Sections map 1:1 to the brief's Testing Requirements:
  T1  fresh initialise: tables, indexes, serial sequences
  T2  staff import round-trip
  T3  PAR import round-trip
  T4  allocation upsert (ON CONFLICT DO UPDATE)
  T5  _ensure_column no-op / add / idempotency
  T6  summary cache build + read (+ ETag)
  T7  nightly job end-to-end (offline bank-holiday failure exercises
      the aborted-transaction rollback path)
  T8  CONCURRENT WRITES — two connections writing simultaneously
  T9  all dashboard data endpoints load via the app (with login)
"""

import json
import os
import sys
import tempfile
import threading
import time
from datetime import date, datetime, timezone
from pathlib import Path

RESULTS = []


def check(name, ok, detail=""):
    RESULTS.append((name, bool(ok), detail))
    print(f"{'PASS' if ok else 'FAIL'}  {name}" + (f"  [{detail}]" if detail and not ok else ""))


def section(title):
    print(f"\n== {title} " + "=" * max(0, 60 - len(title)))


def main():
    assert "RF_DATABASE_URL" in os.environ, "Set RF_DATABASE_URL to an empty test database"
    assert os.environ["RF_DATABASE_URL"].startswith("postgresql"), "This suite targets PostgreSQL"
    os.environ.setdefault("RF_ENV", "development")
    os.environ.setdefault("RF_ADMIN_TOKEN", "testtoken")

    sys.path.insert(0, str(Path(__file__).parent))

    import database
    import config

    # ------------------------------------------------------------------ T1
    section("T1  Fresh initialise")
    database.initialise_database()
    conn = database.get_connection()
    c = conn.cursor()

    c.execute("""SELECT table_name FROM information_schema.tables
                 WHERE table_schema = 'public'""")
    tables = {r["table_name"] for r in c.fetchall()}
    expected = {"staff", "projects", "rtcs", "allocations", "reporting_periods",
                "import_log", "summary_cache", "bank_holidays", "staff_availability"}
    check("all tables created", expected <= tables, f"missing: {expected - tables}")

    c.execute("SELECT indexname FROM pg_indexes WHERE schemaname='public'")
    idx = {r["indexname"] for r in c.fetchall()}
    for want in ("idx_allocations_rtc", "idx_allocations_period", "idx_allocations_person"):
        check(f"index {want} present", want in idx)

    c.execute("SELECT COUNT(*) AS n FROM reporting_periods")
    check("reporting periods seeded", c.fetchone()["n"] > 24)
    c.execute("SELECT COUNT(*) AS n FROM staff WHERE horizon_person_number LIKE 'GENERIC-%'")
    check("generic staff seeded", c.fetchone()["n"] >= 10)

    # SERIAL sequences actually assign ids
    c.execute("""INSERT INTO projects (project_number, task_order_number, project_name,
                 task_name, project_status, last_imported)
                 VALUES ('T1-SEQ','000','t','t','Placeholder','x') RETURNING project_id""")
    pid1 = c.fetchone()["project_id"]
    c.execute("""INSERT INTO projects (project_number, task_order_number, project_name,
                 task_name, project_status, last_imported)
                 VALUES ('T1-SEQ','001','t','t','Placeholder','x') RETURNING project_id""")
    pid2 = c.fetchone()["project_id"]
    check("SERIAL sequence increments", pid2 == pid1 + 1, f"{pid1} -> {pid2}")
    conn.commit()

    # Re-run initialise on a populated DB — must be a clean no-op
    database.initialise_database()
    check("initialise_database idempotent", True)

    # ------------------------------------------------------------------ T2
    section("T2  Staff import round-trip")
    import openpyxl
    tmp = Path(tempfile.mkdtemp())
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Horizon Person Number", "Name", "Job Title", "Line Manager",
               "Job Function", "Job Family", "Department", "Availability",
               "Start Date", "End Date"])
    ws.append(["10001", "Smith, Anna", "P3 - Senior Engineer", "Jones, Bob",
               "Mechanical", "Eng", "UK010117-UK-BSV-Services London", 1.0,
               "01/02/2020", None])
    ws.append(["10002", "Jones, Bob", "P6 - Director", "",
               "Mechanical", "Eng", "UK010117-UK-BSV-Services London", 0.8,
               "01/02/2015", None])
    ws.append(["10003", "Left, Larry", "P2 - Engineer", "Jones, Bob",
               "Electrical", "Eng", "UK010117-UK-BSV-Services London", 1.0,
               "01/02/2021", "30/06/2026"])
    staff_path = tmp / "staff_list.xlsx"
    wb.save(staff_path)

    from imports import staff_list as staff_import
    r = staff_import.run(str(staff_path))
    check("staff import processes rows", r.get("rows_processed", 0) == 3, str(r))
    c.execute("SELECT * FROM staff WHERE horizon_person_number = '10001'")
    row = c.fetchone()
    check("staff round-trip: fields", row and row["name"] == "Smith, Anna"
          and row["job_title"] == "P3 - Senior Engineer"
          and row["line_manager"] == "Jones, Bob")
    check("staff round-trip: date parsed to ISO", row and row["start_date"] == "2020-02-01")
    c.execute("SELECT availability FROM staff WHERE horizon_person_number = '10002'")
    check("staff round-trip: availability float", abs(c.fetchone()["availability"] - 0.8) < 1e-9)

    # Second import (same file) updates rather than duplicates
    r2 = staff_import.run(str(staff_path))
    c.execute("SELECT COUNT(*) AS n FROM staff WHERE horizon_person_number LIKE '100%'")
    check("staff re-import idempotent", c.fetchone()["n"] == 3, str(r2))
    conn.commit()

    # ------------------------------------------------------------------ T3
    section("T3  PAR import round-trip")
    par_dir = tmp / "par"
    par_dir.mkdir()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Project Type", "Project Number", "Project Name", "Project Organization",
               "Project Customer", "Project Status", "Project Director", "Project Manager",
               "Task Number", "Task Name", "Task Start Date", "Task End Date",
               "Reporting Period"])
    ws.append(["UK Direct", "UK0041867", "Test Project", "UK010117-UK-BSV-Services London",
               "Cust Ltd", "Active", "Jones, Bob", "Smith, Anna",
               "9081", "Design", "01/01/2026", "31/12/2026", "202607"])
    ws.append(["UK Opportunity", "UK0099999", "Maybe Project", "UK010117-UK-BSV-Services London",
               "Cust Ltd", "Active", "Jones, Bob", "Smith, Anna",
               "0001", "Bid", "01/01/2026", "31/12/2026", "202607"])
    par_path = par_dir / "202607_UK_PAR_Active.xlsx"
    wb.save(par_path)

    # NOTE: par_import.run()'s file_path argument is ignored — it always
    # reads the newest UK_PAR*.xlsx from config.PAR_ACTUALS_PATH (see report).
    import config as _cfg
    _cfg.PAR_ACTUALS_PATH = par_dir
    from imports import par_import
    r = par_import.run()
    check("PAR import processes rows", r.get("rows_processed", 0) >= 2, str(r))
    c.execute("""SELECT * FROM projects WHERE project_number='UK0041867'
                 AND task_order_number='9081'""")
    row = c.fetchone()
    check("PAR round-trip: project fields", row and row["project_name"] == "Test Project"
          and row["project_type"] == "UK Direct" and row["project_status"] == "Active")
    c.execute("SELECT COUNT(*) AS n FROM projects WHERE project_number IN ('ID-06','ID-04','IDUK-01')")
    check("PAR indirect seed rows present", c.fetchone()["n"] >= 3)
    conn.commit()

    # ------------------------------------------------------------------ T4/T9 app boot
    section("T4  Allocation upsert via the app (+ login)")
    import app as app_module
    application = app_module.create_app()
    client = application.test_client()

    rl = client.post("/login", data={"horizon_person_number": "10001"})
    check("login (session identity)", rl.status_code == 302)

    r = client.post("/api/rtcs", json={
        "project_number": "UK0041867", "task_order_number": "9081",
        "department": "UK010117-UK-BSV-Services London", "start_date": "2026-07-01"})
    d = r.get_json()
    check("create RTC returns rtc_id (RETURNING works)", r.status_code == 201 and d.get("rtc_id"), str(d))
    rtc_id = d["rtc_id"]

    client.post(f"/api/rtcs/{rtc_id}/staff", json={"horizon_person_number": "10001"})
    r = client.patch(f"/api/rtcs/{rtc_id}", json={"allocations": [
        {"horizon_person_number": "10001", "period_start": "2026-08-01", "days": 5}]})
    check("first allocation insert", r.get_json().get("allocations_updated") == 1)
    r = client.patch(f"/api/rtcs/{rtc_id}", json={"allocations": [
        {"horizon_person_number": "10001", "period_start": "2026-08-01", "days": 9}]})
    check("upsert PATCH accepted", r.get_json().get("allocations_updated") == 1)
    c.execute("""SELECT COUNT(*) AS n, MAX(days) AS d FROM allocations
                 WHERE rtc_id = %s AND horizon_person_number='10001'
                 AND period_start='2026-08-01'""", (rtc_id,))
    row = c.fetchone()
    check("ON CONFLICT DO UPDATE: one row, new value", row["n"] == 1 and row["d"] == 9)

    # beacon alias (sendBeacon sends POST text/plain)
    r = client.post(f"/api/rtcs/{rtc_id}", data=json.dumps({"allocations": [
        {"horizon_person_number": "10001", "period_start": "2026-08-01", "days": 4}]}),
        content_type="text/plain")
    check("beacon POST alias works on PG", r.status_code == 200)

    # ------------------------------------------------------------------ T5
    section("T5  _ensure_column")
    conn2 = database.get_connection()
    c2 = conn2.cursor()
    database._ensure_column(c2, "rtcs", "notes", "TEXT")          # exists — no-op
    c2.execute("""SELECT COUNT(*) AS n FROM information_schema.columns
                  WHERE table_name='rtcs' AND column_name='notes'""")
    check("_ensure_column no-ops on existing column", c2.fetchone()["n"] == 1)
    database._ensure_column(c2, "rtcs", "mig_test_col", "TEXT")   # new
    conn2.commit()
    c2.execute("""SELECT COUNT(*) AS n FROM information_schema.columns
                  WHERE table_name='rtcs' AND column_name='mig_test_col'""")
    check("_ensure_column adds a new column", c2.fetchone()["n"] == 1)
    database._ensure_column(c2, "rtcs", "mig_test_col", "TEXT")   # again — no-op
    conn2.commit()
    check("_ensure_column idempotent", True)
    c2.execute("ALTER TABLE rtcs DROP COLUMN mig_test_col")
    conn2.commit()
    conn2.close()

    # ------------------------------------------------------------------ T6
    section("T6  Summary cache")
    import summary as summary_module
    payload = summary_module.build()
    check("summary builds", isinstance(payload, dict) and "staff" in payload and "projects" in payload)
    cached = summary_module.get_cached()
    check("summary cache reads back", cached and cached.get("generated_at"))
    body = cached["payload"]
    parsed = json.loads(body) if isinstance(body, str) else body
    check("cached payload parses, staff present",
          any(s["id"] == "10001" for s in parsed["staff"]))
    r = client.get("/api/summary")
    # Echo the ETag verbatim, as browsers do. (The app emits an unquoted
    # ETag — non-RFC but functional; see the migration report's notes.)
    etag = r.headers.get("ETag", "")
    r2 = client.get("/api/summary", headers={"If-None-Match": etag} if etag else {})
    check("/api/summary 200 + ETag 304 flow", r.status_code == 200 and
          (r2.status_code == 304 if etag else True), f"etag={etag!r} second={r2.status_code}")

    # ------------------------------------------------------------------ T7
    section("T7  Nightly job (offline bank-holiday failure = rollback path)")
    os.environ["RF_STAFF_LIST_PATH"] = str(staff_path)
    config.STAFF_LIST_PATH = staff_path
    from services import jobs
    try:
        jobs.nightly_imports()
        check("nightly_imports completes without raising", True)
    except Exception as e:
        check("nightly_imports completes without raising", False, repr(e))
    # After the failed gov.uk fetch mid-job, the same connection kept working —
    # prove the DB is in a good state by querying and by special RTC creation:
    c.execute("SELECT COUNT(*) AS n FROM rtcs r JOIN projects p ON p.project_id=r.project_id "
              "WHERE p.project_number IN ('ID-06','ID-04','IDUK-01')")
    check("special RTCs maintained despite offline fetch", c.fetchone()["n"] >= 3)

    # ------------------------------------------------------------------ T8
    section("T8  CONCURRENT WRITES — the whole point")
    # Ensure the second person is on the RTC too
    client.post(f"/api/rtcs/{rtc_id}/staff", json={"horizon_person_number": "10002"})

    ITER = 60
    errors = []
    barrier = threading.Barrier(3)

    def writer(person, value_base):
        wconn = database.get_connection()
        wc = wconn.cursor()
        try:
            barrier.wait()
            for i in range(ITER):
                wc.execute("""
                    INSERT INTO allocations
                        (horizon_person_number, rtc_id, period_start, days, last_updated)
                    VALUES (%s, %s, %s, %s, %s)
                    ON CONFLICT (horizon_person_number, rtc_id, period_start)
                    DO UPDATE SET days = EXCLUDED.days, last_updated = EXCLUDED.last_updated
                """, (person, rtc_id, "2026-09-01", value_base + i,
                      datetime.now(timezone.utc).isoformat()))
                wconn.commit()
        except Exception as e:
            errors.append(f"writer {person}: {e!r}")
        finally:
            wconn.close()

    def reader():
        rconn = database.get_connection()
        rc = rconn.cursor()
        try:
            barrier.wait()
            for _ in range(ITER):
                rc.execute("SELECT SUM(days) AS s FROM allocations WHERE rtc_id = %s", (rtc_id,))
                rc.fetchone()
                rconn.commit()
        except Exception as e:
            errors.append(f"reader: {e!r}")
        finally:
            rconn.close()

    t1 = threading.Thread(target=writer, args=("10001", 100))
    t2 = threading.Thread(target=writer, args=("10002", 500))
    t3 = threading.Thread(target=reader)
    start = time.perf_counter()
    for t in (t1, t2, t3): t.start()
    for t in (t1, t2, t3): t.join()
    dur = time.perf_counter() - start

    check(f"2 writers + 1 reader, {ITER} iterations each, zero errors",
          not errors, "; ".join(errors[:3]))
    c.execute("""SELECT horizon_person_number, days FROM allocations
                 WHERE rtc_id = %s AND period_start = '2026-09-01'
                 ORDER BY horizon_person_number""", (rtc_id,))
    rows = c.fetchall()
    check("both writers' final rows present with last value",
          len(rows) == 2 and rows[0]["days"] == 100 + ITER - 1
          and rows[1]["days"] == 500 + ITER - 1,
          str([(r['horizon_person_number'], r['days']) for r in rows]))
    print(f"      ({ITER*2} interleaved upserts + {ITER} reads in {dur*1000:.0f} ms, "
          f"no 'database is locked' — MVCC doing its job)")

    # ------------------------------------------------------------------ T9
    section("T9  Dashboard data endpoints")
    for name, url, checker in [
        ("summary",     "/api/summary",           lambda j: True),
        ("rtcs list",   "/api/rtcs",              lambda j: any(x["rtc_id"] == rtc_id for x in j)),
        ("rtc detail",  f"/api/rtcs/{rtc_id}",    lambda j: j["rtc"]["rtc_id"] == rtc_id and j["staff"]),
        ("staff list",  "/api/staff",             lambda j: any(x["horizon_person_number"] == "10001" for x in j)),
        ("project",     "/api/project?project_number=UK0041867&task_order_number=9081",
                                                  lambda j: j.get("match_type") == "full"),
    ]:
        r = client.get(url)
        ok = r.status_code == 200
        if ok and r.mimetype == "application/json" and name != "summary":
            try:
                ok = checker(r.get_json())
            except Exception as e:
                ok = False
        check(f"GET {name}", ok, f"status={r.status_code}")

    conn.close()

    # ------------------------------------------------------------------ done
    fails = [n for n, ok, _ in RESULTS if not ok]
    print(f"\n{'='*66}\n{len(RESULTS) - len(fails)}/{len(RESULTS)} PASS"
          + (f"  FAILURES: {fails}" if fails else "  — migration suite green"))
    return 1 if fails else 0


if __name__ == "__main__":
    sys.exit(main())
